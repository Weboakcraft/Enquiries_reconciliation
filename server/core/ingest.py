"""
Ingestion: read a workbook, work out which export format it is, and turn its rows
into canonical Lead records.

Nothing about the three known formats is hard-coded here. Everything the engine
knows about a format lives in config/sources.json, so a new export can be
supported by adding a profile there.
"""
from __future__ import annotations

import datetime as dt
import re
from typing import Any

import openpyxl

MAX_HEADER_SCAN = 8       # how many top rows to try as the header row
MAX_PROBE_COLS = 80       # ignore the long tail of empty columns Excel leaves behind


# --------------------------------------------------------------------------- utils
def norm_key(s: Any) -> str:
    """Loose column-name key: case, spaces and punctuation don't matter."""
    return re.sub(r"[^a-z0-9]", "", str(s).lower()) if s is not None else ""


def clean(v: Any, null_tokens: list[str]) -> str:
    if v is None:
        return ""
    s = re.sub(r"\s+", " ", str(v)).strip()
    return "" if s.lower() in null_tokens else s


def to_int(v: Any) -> int:
    try:
        n = int(float(str(v).strip()))
    except (TypeError, ValueError):
        return 0
    return n if 0 < n < 1_000_000 else 0


def parse_date(v: Any, formats: list[str]) -> dt.date | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for f in formats + ["%d-%m-%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S",
                        "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"]:
        try:
            return dt.datetime.strptime(s, f).date()
        except ValueError:
            continue
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return dt.date(*map(int, m.groups()))
    return None


# --------------------------------------------------------------------------- detection
def _header_candidates(ws) -> list[tuple[int, list[Any]]]:
    out = []
    for r in range(1, min(MAX_HEADER_SCAN, ws.max_row) + 1):
        row = [c.value for c in next(ws.iter_rows(min_row=r, max_row=r, max_col=MAX_PROBE_COLS))]
        if sum(1 for v in row if isinstance(v, str) and v.strip()) >= 3:
            out.append((r, row))
    return out


def score_profile(profile: dict, headers: list[Any]) -> float:
    """1.0 = every required column present and at least one 'any' column."""
    keys = {norm_key(h) for h in headers if h is not None}
    req = [norm_key(c) for c in profile["match"]["required"]]
    anyc = [norm_key(c) for c in profile["match"].get("any", [])]
    if not req or not all(k in keys for k in req):
        return 0.0
    if anyc and not any(k in keys for k in anyc):
        return 0.5
    return 1.0


def detect_sheets(path: str, cfg: dict) -> list[dict]:
    """Return one detection record per sheet that matched a profile (plus misses)."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    results = []
    for ws in wb.worksheets:
        best = None
        for header_row, headers in _header_candidates(ws):
            for prof in cfg["profiles"]:
                s = score_profile(prof, headers)
                if s >= 1.0 and (best is None or s > best["score"]):
                    best = {"sheet": ws.title, "profile": prof, "header_row": header_row,
                            "headers": [h for h in headers if h is not None], "score": s}
            if best:
                break
        if best:
            results.append(best)
        else:
            first = _header_candidates(ws)
            results.append({"sheet": ws.title, "profile": None, "header_row": None,
                            "headers": [h for h in (first[0][1] if first else []) if h is not None],
                            "score": 0.0})
    wb.close()
    return results


# --------------------------------------------------------------------------- field resolution
class RowReader:
    """Resolves configured field names against this sheet's actual header row."""

    def __init__(self, headers: list[Any]):
        self.index: dict[str, int] = {}
        for i, h in enumerate(headers):
            k = norm_key(h)
            if k and k not in self.index:
                self.index[k] = i
        self.headers = headers

    def get(self, row: tuple, names: list[str]) -> Any:
        for n in names:
            i = self.index.get(norm_key(n))
            if i is not None and i < len(row):
                v = row[i]
                if v is not None and str(v).strip() != "":
                    return v
        return None

    def has(self, name: str) -> bool:
        return norm_key(name) in self.index


def map_quality(raw: str, cfg: dict, blank: str) -> str:
    s = (raw or "").lower()
    if not s:
        return blank
    for rule in cfg["quality_rules"]:
        if rule["contains"] in s:
            return rule["value"]
    return blank


def map_stage(raw: str, stage_cfg: dict, canonical: list[str]) -> tuple[str, bool]:
    """Returns (stage, was_recognised). Unrecognised non-blank values are flagged."""
    s = (raw or "").strip()
    if not s:
        return stage_cfg.get("blank", "New"), True
    m = {k.upper(): v for k, v in stage_cfg.get("map", {}).items()}
    if s.upper() in m:
        return m[s.upper()], True
    for c in canonical:                       # already canonical?
        if s.lower() == c.lower():
            return c, True
    return stage_cfg.get("blank", "New"), False


def resolve_rep(raw: Any, cfg: dict) -> str:
    s = re.sub(r"\s+", " ", str(raw or "")).strip()
    if not s:
        return cfg["canonical"]["unassigned_label"]
    alias = cfg.get("rep_aliases", {})
    hit = alias.get(s.lower())
    return hit if hit else s.title()


def classify_disposition(note: str, quality: str, stage: str, cfg: dict) -> dict:
    d = re.sub(r"\s+", " ", (note or "")).strip().lower()
    q = (quality or "").lower()
    for rule in cfg["disposition_rules"]:
        if "all_regex" in rule:
            if all(re.search(p, d) for p in rule["all_regex"]):
                return {"id": rule["id"], "label": rule["label"], "severity": rule.get("severity", "normal")}
        elif re.search(rule["any_regex"], d):
            return {"id": rule["id"], "label": rule["label"], "severity": rule.get("severity", "normal")}
    if not d:
        if "interest" in q:
            return {"id": "no_req", "label": "No requirement", "severity": "normal"}
        if "pick" in q:
            return {"id": "no_resp", "label": "No response / call not picked", "severity": "normal"}
    if stage in cfg["canonical"]["closed_no_order"]:
        return {"id": "not_logged", "label": "Reason not logged", "severity": "gap"}
    return {"id": "none", "label": "Not applicable", "severity": "normal"}


# --------------------------------------------------------------------------- main entry
def read_file(path: str, filename: str, cfg: dict) -> dict:
    """Read one workbook -> {leads, sheets, rows_read, rows_blank, unmapped_stage_values}."""
    nulls = [t.lower() for t in cfg["null_tokens"]]
    canonical = cfg["canonical"]["stages"]
    detections = detect_sheets(path, cfg)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    leads, sheets, unmapped = [], [], []
    rows_read = rows_blank = 0

    for det in detections:
        if det["profile"] is None:
            sheets.append({"file": filename, "sheet": det["sheet"], "profile": None,
                           "label": "not recognised", "rows": 0,
                           "headers": [str(h) for h in det["headers"][:25]]})
            continue

        prof, ws = det["profile"], wb[det["sheet"]]
        rr = RowReader(det["headers"])
        f = prof["fields"]
        sn = prof.get("source_name", {})
        n_sheet = 0

        for excel_row, row in enumerate(
            ws.iter_rows(min_row=det["header_row"] + 1, max_col=len(det["headers"]), values_only=True),
            start=det["header_row"] + 1,
        ):
            rows_read += 1
            if not any(c is not None and str(c).strip() != "" for c in row):
                rows_blank += 1
                continue

            # ---- source name
            if "const" in sn:
                source = sn["const"]
                if sn.get("sub_column") and rr.has(sn["sub_column"]):
                    sub = clean(rr.get(row, [sn["sub_column"]]), nulls).lower()
                    source_detail = sn.get("sub_map", {}).get(sub, sub.title() or source)
                else:
                    source_detail = source
            else:
                raw_src = clean(rr.get(row, [sn.get("column", "")]), nulls)
                source = sn.get("default", "Unknown")
                for frag, name in sn.get("normalize", {}).items():
                    if frag in raw_src.lower():
                        source = name
                        break
                source_detail = source

            raw_stage = clean(rr.get(row, prof["stage"]["from"]), nulls)
            stage, ok = map_stage(raw_stage, prof["stage"], canonical)
            if not ok:
                unmapped.append({"file": filename, "sheet": det["sheet"], "row": excel_row,
                                 "value": raw_stage, "field": prof["stage"]["from"][0]})

            raw_quality = clean(rr.get(row, prof["quality"]["from"]), nulls)
            quality = map_quality(raw_quality, cfg, prof["quality"].get("blank", "Not Rated"))
            note = clean(rr.get(row, f["note"]["from"]), nulls)
            qty_raw = rr.get(row, f["qty"]["from"])
            qty = to_int(qty_raw)
            # a numeric field holding the text "null" is a real defect worth surfacing
            qty_defect = bool(f["qty"]["from"]) and qty == 0 and qty_raw is not None \
                and str(qty_raw).strip().lower() in ("null", "none", "n/a", "-")

            date = parse_date(rr.get(row, f["date"]["from"]), f["date"].get("formats", []))
            rep = resolve_rep(rr.get(row, f["rep"]["from"]), cfg)

            leads.append({
                "src_file": filename, "src_sheet": det["sheet"], "src_row": excel_row,
                "profile": prof["id"],
                "source": source, "source_detail": source_detail,
                "date": date.isoformat() if date else None,
                "rep": rep,
                "assigned": rep != cfg["canonical"]["unassigned_label"],
                "customer": clean(rr.get(row, f["customer"]["from"]), nulls),
                "phone": re.sub(r"\D", "", str(rr.get(row, f["phone"]["from"]) or ""))[-10:],
                "city": clean(rr.get(row, f["city"]["from"]), nulls),
                "product": clean(rr.get(row, f["product"]["from"]), nulls)[:80],
                "qty": qty, "qty_defect": qty_defect,
                "stage": stage, "raw_stage": raw_stage,
                "quality": quality, "raw_quality": raw_quality,
                "note": note,
                "disposition": classify_disposition(note, quality, stage, cfg),
            })
            n_sheet += 1

        sheets.append({"file": filename, "sheet": det["sheet"], "profile": prof["id"],
                       "label": prof["label"], "rows": n_sheet,
                       "headers": [str(h) for h in det["headers"][:25]]})

    wb.close()
    return {"leads": leads, "sheets": sheets, "rows_read": rows_read,
            "rows_blank": rows_blank, "unmapped_stage_values": unmapped}
