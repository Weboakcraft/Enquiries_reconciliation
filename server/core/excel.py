"""Writes the master workbook: master report, reconciliation, analysis and raw data."""
from __future__ import annotations

import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

BAR, NAVY, PINK, PINKT = "2E75B6", "1F4E79", "FFC7CE", "9C0006"
CREAM, CREAMT, GREY = "FFF2CC", "C55A11", "6B7684"
GREEN, GREENL, RED, REDL = "1E7A4B", "DCF0E4", "B3261E", "FBE3E1"
AMBER, AMBERL, A = "B26B00", "FDF0DC", "Arial"

F_T = Font(name=A, size=16, bold=True, color=NAVY)
F_ST = Font(name=A, size=9, color=GREY, italic=True)
F_BAR = Font(name=A, size=11, bold=True, color="FFFFFF")
F_H = Font(name=A, size=10, bold=True, color="FFFFFF")
F_B = Font(name=A, size=10)
F_BB = Font(name=A, size=10, bold=True)
F_SEC = Font(name=A, size=11, bold=True, color=NAVY)
F_N = Font(name=A, size=9, color=GREY)
F_PINK = Font(name=A, size=10, bold=True, color=PINKT)
F_TOT = Font(name=A, size=10, bold=True, color=CREAMT)
FL_BAR, FL_H = PatternFill("solid", fgColor=BAR), PatternFill("solid", fgColor=NAVY)
FL_PINK, FL_CREAM = PatternFill("solid", fgColor=PINK), PatternFill("solid", fgColor=CREAM)
FL_G, FL_R = PatternFill("solid", fgColor=GREENL), PatternFill("solid", fgColor=REDL)
FL_A = PatternFill("solid", fgColor=AMBERL)
_t = Side(style="thin", color="9EB6CE")
BORD = Border(left=_t, right=_t, top=_t, bottom=_t)
C = Alignment(horizontal="center", vertical="center")
L = Alignment(horizontal="left", vertical="center")
W = Alignment(horizontal="center", vertical="center", wrap_text=True)
INT, PCT, PCT2, PLAIN = '#,##0;-#,##0;"-"', '0.0%;-0.0%;"-"', '0.00%;-0.00%;"-"', "#,##0"


def _title(ws, t, sub, width):
    ws["A1"], ws["A2"] = t, sub
    ws["A1"].font, ws["A2"].font = F_T, F_ST
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 24, 14
    ws.sheet_view.showGridLines = False


def _bar(ws, row, text, width):
    c = ws.cell(row=row, column=1, value=text)
    c.font, c.fill, c.alignment = F_BAR, FL_BAR, L
    for i in range(2, width + 1):
        ws.cell(row=row, column=i).fill = FL_BAR
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    ws.row_dimensions[row].height = 20


def _hdr(ws, row, vals, h=30):
    for i, v in enumerate(vals):
        c = ws.cell(row=row, column=1 + i, value=v)
        c.font, c.fill, c.alignment, c.border = F_H, FL_H, W, BORD
    ws.row_dimensions[row].height = h


def _put(ws, r, c, v, fmt=None, font=F_B, align=C, fill=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font, cell.alignment, cell.border = font, align, BORD
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    return cell


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[CL(i)].width = w


def build(rec: dict, m: dict, insights: list[dict], cfg: dict, path: str) -> str:
    C_ = cfg["canonical"]
    stages, quals = m["stages"], m["qualities"]
    reps, sources = m["reps"], m["sources"]
    per = f'{m["period"]["start"]} to {m["period"]["end"]}'
    wb = Workbook()

    # ---------------------------------------------------------------- MASTER REPORT
    ws = wb.active
    ws.title = "MASTER REPORT"
    _title(ws, "Master Lead Report",
           f'All sources · {per} · {m["overall"]["total"]} leads · '
           f'generated {dt.date.today():%d %b %Y}', 7 + len(stages))

    _bar(ws, 4, f"  EMPLOYEE WISE SNAPSHOT  (ALL SOURCES, {per})", 4 + len(sources) + len(stages))
    _hdr(ws, 5, ["Salesperson", "Total Enquiry"] + sources + stages + ["Conversion %", "Loss %"])
    r = 6
    for rep in reps:
        un = rep == C_["unassigned_label"]
        fnt, fl = (F_PINK, FL_PINK) if un else (F_B, None)
        b = m["by_rep"][rep]
        _put(ws, r, 1, rep, None, F_PINK if un else F_BB, L, fl)
        _put(ws, r, 2, b["total"], INT, F_PINK if un else F_BB, C, fl)
        for i, s in enumerate(sources):
            _put(ws, r, 3 + i, m["cube"][rep][s]["total"], INT, fnt, C, fl)
        for i, s in enumerate(stages):
            _put(ws, r, 3 + len(sources) + i, b[s], INT, fnt, C, fl)
        c0 = 3 + len(sources) + len(stages)
        _put(ws, r, c0, b["conversion"] / 100, PCT, F_PINK if un else F_BB, C, fl or FL_G)
        _put(ws, r, c0 + 1, b["loss_rate"] / 100, PCT, fnt, C, fl)
        if not un:
            ws.cell(row=r, column=3 + len(sources)).fill = FL_A
        r += 1
    o = m["overall"]
    _put(ws, r, 1, "TOTAL", None, F_TOT, L, FL_CREAM)
    _put(ws, r, 2, o["total"], INT, F_TOT, C, FL_CREAM)
    for i, s in enumerate(sources):
        _put(ws, r, 3 + i, m["by_source"][s]["total"], INT, F_TOT, C, FL_CREAM)
    for i, s in enumerate(stages):
        _put(ws, r, 3 + len(sources) + i, o[s], INT, F_TOT, C, FL_CREAM)
    c0 = 3 + len(sources) + len(stages)
    _put(ws, r, c0, o["conversion"] / 100, PCT, F_TOT, C, FL_CREAM)
    _put(ws, r, c0 + 1, o["loss_rate"] / 100, PCT, F_TOT, C, FL_CREAM)
    end1 = r

    b2 = end1 + 3
    _bar(ws, b2, "  SOURCE WISE SNAPSHOT", 10)
    _hdr(ws, b2 + 1, ["Source", "Total Enquiry", "% Share", "Assigned", "Unassigned",
                      "Won", "Closed no order", "Qualified", "Hot Leads", "Qty (pcs)"])
    r = b2 + 2
    for s in sources:
        b = m["by_source"][s]
        _put(ws, r, 1, s, None, F_BB, L)
        _put(ws, r, 2, b["total"], INT, F_BB)
        _put(ws, r, 3, b["total"] / o["total"] if o["total"] else 0, PCT)
        _put(ws, r, 4, b["total"] - b["unassigned"], INT)
        _put(ws, r, 5, b["unassigned"], INT, F_PINK if b["unassigned"] else F_B, C,
             FL_PINK if b["unassigned"] else None)
        _put(ws, r, 6, b["Won"], INT, F_B, C, FL_G)
        _put(ws, r, 7, b["closed_no_order"], INT, F_B, C, FL_R)
        _put(ws, r, 8, b["Qualified"], INT)
        _put(ws, r, 9, b["hot"], INT, F_BB)
        _put(ws, r, 10, b["qty"], PLAIN, F_BB)
        r += 1
    _put(ws, r, 1, "TOTAL", None, F_TOT, L, FL_CREAM)
    for col, v in [(2, o["total"]), (4, o["total"] - o["unassigned"]), (5, o["unassigned"]),
                   (6, o["Won"]), (7, o["closed_no_order"]), (8, o["Qualified"]),
                   (9, o["hot"]), (10, o["qty"])]:
        _put(ws, r, col, v, PLAIN if col == 10 else INT, F_TOT, C, FL_CREAM)
    _put(ws, r, 3, 1, PCT, F_TOT, C, FL_CREAM)
    end2 = r

    b3 = end2 + 2
    _bar(ws, b3, "  STAGE WISE SNAPSHOT  (UNIFIED)", 3 + len(sources))
    _hdr(ws, b3 + 1, ["Stage (Unified)", "No. of Leads", "% Share"] + sources)
    r = b3 + 2
    for row in m["stage_matrix"]:
        s = row["stage"]
        fl = FL_G if s == "Won" else FL_R if s in C_["closed_no_order"] else \
            FL_A if s == C_["untouched_stage"] else None
        _put(ws, r, 1, s, None, F_BB, L, fl)
        _put(ws, r, 2, row["total"], INT, F_BB, C, fl)
        _put(ws, r, 3, row["share"] / 100, PCT, F_B, C, fl)
        for i, src in enumerate(sources):
            _put(ws, r, 4 + i, row[src], INT, F_B, C, fl)
        r += 1
    _put(ws, r, 1, "TOTAL", None, F_TOT, L, FL_CREAM)
    _put(ws, r, 2, o["total"], INT, F_TOT, C, FL_CREAM)
    _put(ws, r, 3, 1, PCT, F_TOT, C, FL_CREAM)
    for i, src in enumerate(sources):
        _put(ws, r, 4 + i, m["by_source"][src]["total"], INT, F_TOT, C, FL_CREAM)
    end3 = r

    b4 = end3 + 2
    _bar(ws, b4, "  LEAD QUALITY SNAPSHOT  (HOT / WARM / COLD)", 5 + len(sources))
    _hdr(ws, b4 + 1, ["Lead Quality", "Total", "% Share"] + sources + ["Qualified", "Not Qualified"])
    r = b4 + 2
    for row in m["quality_matrix"]:
        q = row["quality"]
        fl = FL_G if q == "Hot Lead" else FL_R if q in ("Not Interested", "Don't Pick Call") else None
        _put(ws, r, 1, q, None, F_BB, L, fl)
        _put(ws, r, 2, row["total"], INT, F_BB, C, fl)
        _put(ws, r, 3, row["share"] / 100, PCT, F_B, C, fl)
        for i, src in enumerate(sources):
            _put(ws, r, 4 + i, row[src], INT, F_B, C, fl)
        _put(ws, r, 4 + len(sources), row["qualified"], INT, F_B, C, fl)
        _put(ws, r, 5 + len(sources), row["not_qualified"], INT, F_B, C, fl)
        r += 1
    _put(ws, r, 1, "TOTAL", None, F_TOT, L, FL_CREAM)
    _put(ws, r, 2, o["total"], INT, F_TOT, C, FL_CREAM)
    _put(ws, r, 3, 1, PCT, F_TOT, C, FL_CREAM)
    _widths(ws, [22, 13] + [11] * len(sources) + [11] * len(stages) + [12, 10])

    # ---------------------------------------------------------------- RECONCILIATION
    ws = wb.create_sheet("Reconciliation")
    _title(ws, "Reconciliation & Audit Trail",
           "Every row that entered the engine is accounted for below. "
           "If the control total balances, nothing was silently dropped.", 9)
    t = rec["control_totals"]
    _bar(ws, 4, "  CONTROL TOTALS", 9)
    _hdr(ws, 5, ["File", "Sheet(s) detected", "Format", "Rows read", "Blank rows",
                 "Parsed", "No date", "Out of period", "Duplicates", "Accepted"])
    r = 6
    for p in rec["per_file"]:
        _put(ws, r, 1, p["file"], None, F_BB, L)
        _put(ws, r, 2, ", ".join(s["sheet"] for s in p["sheets"]), None, F_B, L)
        _put(ws, r, 3, ", ".join(s["label"] for s in p["sheets"]), None, F_B, L)
        for i, k in enumerate(["rows_read", "rows_blank", "parsed", "no_date",
                               "out_of_period", "duplicates", "accepted"]):
            _put(ws, r, 4 + i, p[k], INT, F_BB if k == "accepted" else F_B)
        r += 1
    _put(ws, r, 1, "TOTAL", None, F_TOT, L, FL_CREAM)
    _put(ws, r, 2, "", None, F_TOT, L, FL_CREAM)
    _put(ws, r, 3, "", None, F_TOT, L, FL_CREAM)
    for i, k in enumerate(["rows_read", "rows_blank", "parsed", "no_date",
                           "out_of_period", "duplicates", "accepted"]):
        _put(ws, r, 4 + i, t[k], INT, F_TOT, C, FL_CREAM)
    r += 2
    ws.cell(row=r, column=1,
            value=("BALANCED — parsed − no-date − out-of-period − duplicates = accepted"
                   if t["balanced"] else "OUT OF BALANCE — investigate before using this report")
            ).font = Font(name=A, size=11, bold=True, color=GREEN if t["balanced"] else RED)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)

    r += 2
    _bar(ws, r, "  EXCEPTION REGISTER", 6)
    _hdr(ws, r + 1, ["Severity", "Type", "What was found", "Detail", "Source row", "Suggested fix"])
    rr = r + 2
    for e in rec["exceptions"]:
        fl = FL_PINK if e["severity"] == "critical" else FL_R if e["severity"] == "high" \
            else FL_A if e["severity"] == "medium" else None
        _put(ws, rr, 1, e["severity"].upper(), None, F_BB, C, fl)
        _put(ws, rr, 2, e["kind"], None, F_B, L, fl)
        _put(ws, rr, 3, e["message"], None, F_B, L, fl)
        _put(ws, rr, 4, e["detail"][:180], None, F_N, L, fl)
        _put(ws, rr, 5, e["ref"], None, F_N, L, fl)
        _put(ws, rr, 6, e["fix"], None, F_N, L, fl)
        rr += 1
    ws.auto_filter.ref = f"A{r + 1}:F{rr - 1}"
    _widths(ws, [13, 18, 46, 52, 34, 44, 13, 13, 13, 12])

    # ---------------------------------------------------------------- INSIGHTS
    ws = wb.create_sheet("Findings & Actions")
    _title(ws, "Findings & Recommended Actions",
           f"Generated from the reconciled data for {per}", 8)
    _hdr(ws, 4, ["#", "Severity", "Finding", "Why it matters", "Action"])
    r = 5
    for i, f in enumerate(insights, 1):
        fl = FL_PINK if f["sev"] == "critical" else FL_R if f["sev"] == "high" else FL_A
        _put(ws, r, 1, i, INT, F_BB, C, fl)
        _put(ws, r, 2, f["sev"].upper(), None, F_BB, C, fl)
        c = _put(ws, r, 3, f["title"], None, F_BB, L, fl)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c = _put(ws, r, 4, f["why"], None, F_N, L, fl)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c = _put(ws, r, 5, f["action"], None, F_B, L, fl)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 34
        r += 1
    _widths(ws, [5, 12, 52, 62, 46])

    # ---------------------------------------------------------------- UNASSIGNED
    if m["unassigned"]:
        ws = wb.create_sheet("Unassigned Leads")
        _title(ws, "Unassigned Leads — no owner",
               f'{len(m["unassigned"])} leads reached the system with a blank salesperson and '
               f'were never picked up. Largest quantity first.', 8)
        _hdr(ws, 4, ["Qty (pcs)", "Age (days)", "Date", "Source", "Customer",
                     "Phone", "City / Location", "Requirement"])
        r = 5
        for u in m["unassigned"]:
            _put(ws, r, 1, u["qty"], PLAIN, F_BB)
            _put(ws, r, 2, u["age"], INT, Font(name=A, size=10, bold=True, color=RED)
                 if u["age"] >= 4 else F_B)
            _put(ws, r, 3, dt.date.fromisoformat(u["date"]), "DD-MMM-YY")
            _put(ws, r, 4, u["source"])
            _put(ws, r, 5, u["customer"], None, F_B, L)
            _put(ws, r, 6, u["phone"])
            _put(ws, r, 7, u["city"], None, F_B, L)
            _put(ws, r, 8, u["product"], None, F_B, L)
            if u["qty"] >= 50:
                for c_ in range(1, 9):
                    ws.cell(row=r, column=c_).fill = FL_PINK
            r += 1
        _widths(ws, [11, 11, 12, 12, 26, 14, 20, 46])

    # ---------------------------------------------------------------- DEMAND / LOSS
    ws = wb.create_sheet("Demand & Loss")
    _title(ws, "Demand Lost & Disposition Analysis",
           "Pieces at stake behind every lead closed without an order.", 6)
    d = m["demand_totals"]
    _bar(ws, 4, "  DEMAND AT A GLANCE (pieces)", 6)
    _hdr(ws, 5, ["Total demand", "Closed without order", "Still open", "Won",
                 "Rejected on product range", "% of demand closed"])
    _put(ws, 6, 1, d["total"], PLAIN, F_BB)
    _put(ws, 6, 2, d["closed"], PLAIN, F_BB, C, FL_R)
    _put(ws, 6, 3, d["open"], PLAIN, F_BB, C, FL_A)
    _put(ws, 6, 4, d["won"], PLAIN, F_BB, C, FL_G)
    _put(ws, 6, 5, d["product_gap"], PLAIN, F_BB, C, FL_A)
    _put(ws, 6, 6, (d["closed"] / d["total"]) if d["total"] else 0, PCT, F_BB)
    _bar(ws, 8, "  WHY THAT DEMAND WAS LOST", 4)
    _hdr(ws, 9, ["Reason", "Leads", "Pieces at stake", "% of lost demand"])
    r, tq = 10, sum(x["qty"] for x in m["demand"]) or 1
    for x in m["demand"]:
        fl = FL_A if x["severity"] in ("product", "recoverable") else None
        _put(ws, r, 1, x["label"], None, F_BB, L, fl)
        _put(ws, r, 2, x["leads"], INT, F_B, C, fl)
        _put(ws, r, 3, x["qty"], PLAIN, F_BB, C, fl)
        _put(ws, r, 4, x["qty"] / tq, PCT, F_B, C, fl)
        r += 1
    r += 1
    _bar(ws, r, "  EVERY DISPOSITION (all leads closed without an order)", 4)
    _hdr(ws, r + 1, ["Reason", "Leads", "% of closed", "Pieces"])
    rr = r + 2
    for x in m["dispositions"]:
        _put(ws, rr, 1, x["label"], None, F_BB, L)
        _put(ws, rr, 2, x["leads"], INT)
        _put(ws, rr, 3, x["share"] / 100, PCT)
        _put(ws, rr, 4, x["qty"], PLAIN)
        rr += 1
    rr += 1
    _bar(ws, rr, "  WORTH CHASING AGAIN", 9)
    _hdr(ws, rr + 1, ["Qty", "Date", "Source", "Salesperson", "Customer",
                      "Requirement", "Quality", "Stage", "Note logged"])
    r2 = rr + 2
    for x in m["recoverable"]:
        _put(ws, r2, 1, x["qty"], PLAIN, F_BB)
        _put(ws, r2, 2, dt.date.fromisoformat(x["date"]), "DD-MMM-YY")
        _put(ws, r2, 3, x["source"])
        _put(ws, r2, 4, x["rep"], None, F_B, L)
        _put(ws, r2, 5, x["customer"], None, F_B, L)
        _put(ws, r2, 6, x["product"], None, F_B, L)
        _put(ws, r2, 7, x["quality"], None, F_B, C,
             FL_G if x["quality"] == "Hot Lead" else None)
        _put(ws, r2, 8, x["stage"], None, F_BB, C, FL_R)
        _put(ws, r2, 9, x["note"] or "— nothing recorded —", None, F_N, L)
        r2 += 1
    _widths(ws, [13, 12, 14, 16, 26, 40, 16, 15, 34])

    # ---------------------------------------------------------------- MASTER DATA
    ws = wb.create_sheet("Master Data")
    _title(ws, "Master Data — reconciled lead register",
           f'{len(rec["leads"])} accepted leads for {per}, normalised across every source.', 17)
    cols = ["Lead ref", "Date", "Source", "Salesperson", "Assigned?", "Customer", "Phone",
            "City", "Product / Requirement", "Qty", "Stage (Unified)", "Original status",
            "Lead Quality", "Note", "Disposition", "Source file", "Row"]
    _hdr(ws, 4, cols)
    r = 5
    order = {s: i for i, s in enumerate(stages)}
    for l in sorted(rec["leads"], key=lambda x: (x["rep"] == C_["unassigned_label"],
                                                 x["rep"], order[x["stage"]], x["date"])):
        vals = [f'{l["source"][:3].upper()}-{l["src_row"]:05d}', dt.date.fromisoformat(l["date"]),
                l["source"], l["rep"], "Yes" if l["assigned"] else "No", l["customer"],
                l["phone"], l["city"], l["product"], l["qty"], l["stage"], l["raw_stage"],
                l["quality"], l["note"], l["disposition"]["label"], l["src_file"], l["src_row"]]
        for i, v in enumerate(vals):
            c = ws.cell(row=r, column=1 + i, value=v)
            c.font, c.border = F_B, BORD
            c.alignment = L if i in (5, 7, 8, 13, 14, 15) else C
            if i == 1:
                c.number_format = "DD-MMM-YY"
            if i == 9:
                c.number_format = PLAIN
        st = l["stage"]
        cell = ws.cell(row=r, column=11)
        stage_fill = FL_G if st == "Won" else FL_R if st in C_["closed_no_order"] \
            else FL_A if st == C_["untouched_stage"] else None
        if stage_fill is not None:
            cell.fill = stage_fill
        cell.font = Font(name=A, size=10, bold=True,
                         color=GREEN if st == "Won" else RED if st in C_["closed_no_order"]
                         else AMBER if st == C_["untouched_stage"] else "000000")
        if not l["assigned"]:
            for c_ in (4, 5):
                ws.cell(row=r, column=c_).fill = FL_PINK
                ws.cell(row=r, column=c_).font = F_PINK
        r += 1
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:Q{r - 1}"
    _widths(ws, [13, 11, 11, 16, 10, 24, 13, 18, 40, 8, 17, 15, 15, 30, 26, 26, 7])

    wb.save(path)
    return path
