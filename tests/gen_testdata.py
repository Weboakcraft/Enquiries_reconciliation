"""Builds three synthetic exports that exercise every branch of the engine."""
import datetime as dt
import random
from openpyxl import Workbook

random.seed(7)
END = dt.date(2026, 9, 20)                      # newest row -> period = 14..20 Sep
DAYS = [END - dt.timedelta(days=i) for i in range(7)][::-1]

REPS = ["Ujala Rajput", "UJALA RAJPUT", "amira saikh", "Amira Sheikh",
        "Mayank Mittal", "Niti Kumari", "  ", "", "Rahul  Verma"]
CITIES = ["Delhi", "Mumbai", "Pune", "Jaipur", "Kolkata", "", "-"]
PRODUCTS = ["Office Chair", "Workstation 4 seater", "Conference Table",
            "Executive Desk", "Visitor Chair", "", "Storage Unit"]
NOTES = [
    "", "call back next week", "not picked, catalogue shared", "fake lead",
    "no requirement", "low price expectation", "requirement of 50 chairs",
    "wants plastec chairs", "delhi base dealer only", "only 1 pec needed",
    "nopt pick call", "catalou shared", "call cut", "not requrement",
    "customer has requir of tables", "godrej brand only", "budget issue",
]
QUALITIES = ["Hot Lead", "Warm", "cold lead", "Don't Pick Call",
             "Not Interested", "", "HOT", "warm lead"]

SHARED_PHONES = ["9811122233", "9822233344", "9833344455"]


def phone(i, dup_pool=False):
    if dup_pool and i % 9 == 0:
        return random.choice(SHARED_PHONES)
    if i % 17 == 0:
        return ""                                  # no phone at all
    if i % 23 == 0:
        return "12345"                             # too short to dedupe on
    return f"9{random.randint(100000000, 999999999)}"


def a_date(i, allow_bad=True):
    if allow_bad and i % 31 == 0:
        return ""                                  # no date -> excluded
    if allow_bad and i % 29 == 0:
        return END - dt.timedelta(days=40)         # out of period
    return DAYS[i % 7]


# ─────────────────────────────────────────── 1. Meta export (text dates)
wb = Workbook()
ws = wb.active
ws.title = "Leads"
ws.append(["created_time_ist", "platform", "full_name", "phone_number",
           "city_state", "lead_status", "assigned_to"])
META_STATUS = ["CREATED", "NEW", "CONTACTED", "QUALIFIED", "WON", "LOST",
               "NOT QUALIFIED", "PROPOSAL", "", "IN REVIEW"]
for i in range(130):
    d = a_date(i)
    ds = "" if d == "" else f"{d.strftime('%d-%m-%Y')} {9 + i % 9:02d}:{i % 60:02d}:00"
    ws.append([ds,
               ["ig", "fb", ""][i % 3],
               ["Jd Buyer", f"Meta Buyer {i}", "buyer", f"Customer {i}"][i % 4],
               phone(i, dup_pool=True),
               CITIES[i % len(CITIES)],
               META_STATUS[i % len(META_STATUS)],
               REPS[i % len(REPS)]])
ws2 = wb.create_sheet("Notes")                     # unrecognised sheet
ws2.append(["Some", "Random", "Columns", "Here"])
ws2.append(["a", "b", "c", "d"])
wb.save("sample_meta.xlsx")

# ─────────────────────────────────────────── 2. IndiaMART workflow (real dates)
wb = Workbook()
ws = wb.active
ws.title = "Sheet1"
ws.append(["Note: exported from IndiaMART", None, None])          # junk row 1
ws.append(["Timestamp", "Buyer Name", "Mobile Number", "City/ Location",
           "Subject", "Quantity", "Assigned Salesperson",
           "Qualification Status", "Lead's Quality", "DISCUSSION"])
IM_STATUS = ["QUALIFIED", "NOT QUALIFIED", "LOST", "WON", "", "PENDING"]
for i in range(160):
    d = a_date(i)
    qty = ["null", 0, 5, 12, 60, 150, "", 1][i % 8]
    ws.append([d if d != "" else None,
               ["India Mart Buyer", f"Buyer {i}", "test", f"Acme {i} Pvt Ltd"][i % 4],
               phone(i + 3, dup_pool=True),
               CITIES[i % len(CITIES)],
               PRODUCTS[i % len(PRODUCTS)],
               qty,
               REPS[(i + 2) % len(REPS)],
               IM_STATUS[i % len(IM_STATUS)],
               QUALITIES[i % len(QUALITIES)],
               NOTES[i % len(NOTES)]])
wb.save("sample_indiamart.xlsx")

# ─────────────────────────────────────────── 3. JustDial workflow (mixed sources)
wb = Workbook()
ws = wb.active
ws.title = "Enquiries"
ws.append(["Enquiry Id", "Date", "Customer", "Mobile", "City", "Product",
           "Qty", "Salesperson", "Stage", "Lead's Quality", "Final Remark's",
           "Source"])
JD_STAGE = ["NEW", "CONTACTED", "QUALIFIED", "QUOTED", "WON", "LOST",
            "NOT QUALIFIED", "", "FOLLOW UP"]
SRC = ["JustDial", "India Mart", "meta ads", "justdial"]
for i in range(150):
    d = a_date(i)
    ws.append([f"E{i:04d}",
               d if d != "" else None,
               ["Jd Buyer", f"JD Cust {i}", "unknown", f"Retail {i}"][i % 4],
               phone(i + 5, dup_pool=True),
               CITIES[i % len(CITIES)],
               PRODUCTS[(i + 1) % len(PRODUCTS)],
               [2, 25, 0, "", 80, 300, 7, "null"][i % 8],
               REPS[(i + 4) % len(REPS)],
               JD_STAGE[i % len(JD_STAGE)],
               QUALITIES[(i + 3) % len(QUALITIES)],
               NOTES[(i + 5) % len(NOTES)],
               SRC[i % len(SRC)]])
wb.save("sample_justdial.xlsx")

print("wrote sample_meta.xlsx, sample_indiamart.xlsx, sample_justdial.xlsx")
